"""Build the dated IAME submission snapshot without changing earlier revisions."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import zipfile

from docx import Document
from docx.shared import Inches, Pt
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'outputs/paper/review_revision'
OUT = ROOT / 'outputs/paper/camera_ready_v1_2026-10-29'
OUT.mkdir(parents=True, exist_ok=True)
PAPER = 'IAME2026_Paper54_Camera_Ready_v1'
RESPONSE = 'IAME2026_Paper54_Response_to_Reviewers_v1'

new_experiment = (
    'A subsequent continuation experiment used 18,199 labelled images from the same public sources, '
    'with 13,772 training and 4,427 validation images. Nimbus recordings were grouped by recording '
    'identifier; the original FDM source was grouped conservatively by capture date, keeping augmented '
    'copies together. On this same validation set, the retained model achieved 51.84% accuracy and '
    '57.62% macro F1; fine-tuning increased these to 83.10% and 84.98%, respectively. Training used a '
    'GTX 1650, batch size four, class weights, eight completed head-training epochs and four completed '
    'fine-tuning epochs, with selection by validation loss. The candidate misclassified 292 good images '
    'as under-extrusion and 252 as stringing. Validation was used for model selection, and overlap with '
    'the retained model\'s earlier training history is unknown. These exploratory results are not an '
    'independent test and must not be compared directly with the 96.52% historical test score. The '
    'candidate has not replaced the retained deployment model.'
)
replacements = {
    'The retained artefacts contain aggregate metrics, the confusion matrix, and model weights, but no paired validation/test probability arrays or original split manifest. ECE, a fitted temperature, and a threshold-performance curve therefore cannot be recovered from the available results alone. No such measurement is claimed here. The implemented 0.60 cutoff remains an unvalidated interface heuristic and must not be read as a statistical safety guarantee. Calibration, temporal-window ablation, and recording-grouped testing remain experimental requirements.':
    'The historical experiment retains aggregate metrics and weights, but lacks paired prediction arrays and its original split manifest. The continuation experiment in Section IV-F retains its split and probabilities, but its validation set was used for model selection. No fitted temperature or independently evaluated calibration score is claimed. The 0.60 cutoff remains an unvalidated interface heuristic. Calibration on separate data, temporal-window ablation, and external recording-grouped testing remain experimental requirements.',
    'The repository retains the model, class order, aggregate metrics, classification report, confusion matrix, training-curve image, and CPU benchmark. These support the reported image-level results. Exact calibration and grouped-split replication additionally require the selected image list, recording identifiers, and per-image probabilities, which are not retained in this package. The public source data must be obtained under their original conditions; a fresh download and random seed alone do not establish an identical historical split.':
    'The historical model, class order, aggregate reports, figures and CPU benchmark are retained, but its original image-level split and prediction arrays are unavailable. The subsequent grouped experiment additionally retains split paths and per-image probabilities. These records describe distinct experiments; a fresh download cannot reconstruct the historical split. Public source data remain subject to their original usage conditions.',
}


def replace_paragraph(paragraph, text):
    properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if properties is not None:
        run._r.insert(0, properties)


paper = Document(SOURCE / 'IAME2026_Revised_Paper.docx')
for paragraph in paper.paragraphs:
    if paragraph.text in replacements:
        replace_paragraph(paragraph, replacements[paragraph.text])
anchor = next(p for p in paper.paragraphs if p.text == 'V. DISCUSSION')
heading_template = next(p for p in paper.paragraphs if p.text.startswith('E. Comparison'))
body_template = next(p for p in paper.paragraphs if p.text.startswith('The table provides context'))
for template, text in [(heading_template, 'F. Exploratory Grouped Continuation Experiment'), (body_template, new_experiment)]:
    element = deepcopy(template._p)
    anchor._p.addprevious(element)
    from docx.text.paragraph import Paragraph
    replace_paragraph(Paragraph(element, anchor._parent), text)
paper.save(OUT / f'{PAPER}.docx')
markdown = (SOURCE / 'IAME2026_Revised_Paper.md').read_text()
for old, new in replacements.items():
    assert old in markdown
    markdown = markdown.replace(old, new)
markdown = markdown.replace('## V. Discussion', '### F. Exploratory Grouped Continuation Experiment\n\n' + new_experiment + '\n\n## V. Discussion')
(OUT / f'{PAPER}.md').write_text(markdown)

response = Document(SOURCE / 'Response_to_Reviewers.docx')
response_md = (SOURCE / 'Response_to_Reviewers.md').read_text()
changes = {}
for paragraph in response.paragraphs:
    text = paragraph.text
    if text.startswith('The repeated general discussion'):
        changes[text] = text + ' Section IV-F now adds an exploratory same-validation-set comparison of the retained model and its locally fine-tuned candidate. The 83.10% result is explicitly separated from the historical 96.52% test result; it does not establish independent generalisation.'
    elif text.startswith('Section V-B adds'):
        changes[text] = text.replace('The retained files lack paired validation/test prediction arrays and the original split manifest,', 'The historical experiment lacks paired validation/test prediction arrays and its original split manifest,').replace('This comment is addressed through', 'The later continuation experiment saves probabilities, but used validation for model selection and has not undergone an independent calibration assessment. This comment is addressed through')
    elif text.startswith('The reported accuracy remains'):
        changes[text] = text.replace('No retraining, new data collection, physical latency experiment, same-split backbone benchmark, calibration fit, or smoothing ablation was performed for this editorial revision.', 'Section IV-F additionally reports completed local continuation training on 13,772 images with grouped validation on 4,427 images. Accuracy increased from 51.84% to 83.10% on that same validation set. This is model-selection evidence with unknown overlap in the source model training history, not a new independent test. No new physical data collection, physical latency experiment, same-split backbone benchmark, calibration fit, or smoothing ablation is claimed.')
for paragraph in response.paragraphs:
    if paragraph.text in changes:
        old = paragraph.text
        replace_paragraph(paragraph, changes[old])
for old, new in changes.items():
    assert old in response_md
    response_md = response_md.replace(old, new)
response.save(OUT / f'{RESPONSE}.docx')
for section in response.sections:
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.8)
for paragraph in response.paragraphs:
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1
    if paragraph.style.name == 'Title':
        for run in paragraph.runs:
            run.font.size = Pt(20)
    elif not paragraph.style.name.startswith('Heading'):
        for run in paragraph.runs:
            run.font.size = Pt(10.5)
response.save(OUT / f'{RESPONSE}.docx')
(OUT / f'{RESPONSE}.md').write_text(response_md)

profile = OUT / '.libreoffice-profile'
subprocess.run(['libreoffice', f'-env:UserInstallation={profile.as_uri()}', '--headless', '--convert-to', 'pdf', '--outdir', str(OUT), str(OUT / f'{PAPER}.docx'), str(OUT / f'{RESPONSE}.docx')], check=True, timeout=120)

pdf_path = OUT / f'{PAPER}.pdf'
with pymupdf.open(pdf_path) as pdf:
    assert len(pdf) <= 6, f'Paper is {len(pdf)} pages; adjust content before delivery'
    count = 0
    for page in pdf:
        for info in page.get_image_info(xrefs=True):
            if info['width'] < 500:
                continue
            name = 'figure1_confusion_matrix.pdf' if info['width'] == info['height'] else 'figure2_class_performance.pdf'
            rect = pymupdf.Rect(info['bbox'])
            with pymupdf.open(SOURCE / name) as vector:
                page.add_redact_annot(rect, fill=(1,1,1))
                page.apply_redactions(images=2, graphics=0, text=1)
                page.show_pdf_page(rect, vector, 0)
            count += 1
    assert count == 2, f'Expected 2 figures, found {count}'
    temp = OUT / 'vector_export.pdf'
    pdf.save(temp, garbage=4, deflate=True)
temp.replace(pdf_path)

checks = {}
for name in [PAPER, RESPONSE]:
    with pymupdf.open(OUT / f'{name}.pdf') as pdf:
        text = '\n'.join(page.get_text() for page in pdf)
        for required in ['96.52%', '83.10%', '4,427']:
            assert required in text, (name, required)
        assert not pdf.is_encrypted
        assert all(abs(p.rect.width-595.28) < 2 and abs(p.rect.height-841.89) < 2 for p in pdf)
        checks[name] = {'pages':len(pdf), 'size':'A4', 'sha256':hashlib.sha256((OUT / f'{name}.pdf').read_bytes()).hexdigest()}
        for i, page in enumerate(pdf):
            page.get_pixmap(matrix=pymupdf.Matrix(1,1)).save(OUT / f'{name}_preview_{i+1}.png')
        (OUT / f'{name}_extracted.txt').write_text(text)
(OUT / 'validation.json').write_text(json.dumps(checks, indent=2))
readme = f'''IAME 2026 - Paper 54 - Camera-ready snapshot v1
Created: 29 October 2026

MAIN MANUSCRIPT: {PAPER}.pdf ({checks[PAPER]['pages']} A4 pages)
REVIEWER RESPONSE: {RESPONSE}.pdf ({checks[RESPONSE]['pages']} A4 pages)

Upload the main manuscript as the paper. Upload the response only in the field
provided/requested by CMT. Do not append the response to the six-page manuscript.
This package has not been submitted to CMT.

Editable DOCX and Markdown sources are included for the next iteration.
The manuscript retains the historical 96.52% test result and separately reports
83.10% grouped selection-validation accuracy from local continuation training.
The latter is not an independent test. No 95% accuracy or calibration is invented.
Both manuscript figures derive from retained historical project results and are
embedded as vectors. They do not depict the new candidate's confusion matrix.

Before submitting: supervisors must approve names/affiliations/content; perform
the required institution-approved similarity check and confirm the individual
CMT deadline, forms and registration requirements. No verified similarity report
is included. Do not describe this package as conference-approved.

For later iterations, create v2 rather than silently replacing these files.
Build script: tools/build_camera_ready_v1.py
'''
(OUT / 'READ_ME.txt').write_text(readme)
figures = OUT / 'figures'
figures.mkdir(exist_ok=True)
for stem in ['figure1_confusion_matrix', 'figure2_class_performance']:
    for extension in ['pdf', 'svg', 'png']:
        shutil.copy2(SOURCE / f'{stem}.{extension}', figures / f'{stem}.{extension}')
with zipfile.ZipFile(OUT / 'IAME2026_Paper54_Camera_Ready_v1.zip', 'w', zipfile.ZIP_DEFLATED) as archive:
    for path in sorted(OUT.iterdir()):
        if path.is_file() and (path.suffix in {'.pdf','.docx','.md'} or path.name in {'READ_ME.txt','validation.json'}):
            archive.write(path, path.name)
    for path in sorted(figures.iterdir()):
        archive.write(path, 'figures/' + path.name)
print(json.dumps(checks, indent=2))
print(OUT)
