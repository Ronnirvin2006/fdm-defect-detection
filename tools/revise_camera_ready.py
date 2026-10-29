"""Prepare the review revision from the retained manuscript and measured results."""
import importlib.util
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs/paper/review_revision"
OUT.mkdir(parents=True, exist_ok=True)
source = (ROOT / "outputs/paper/FDM_Defect_Detection_Conference_Paper.md").read_text()


def section(start, end, replacement):
    global source
    a, b = source.index(start), source.index(end)
    source = source[:a] + replacement.strip() + "\n\n" + source[b:]


source = source.replace("five-frame probability averaging", "five-prediction probability averaging")
source = source.replace("five-frame averaging", "five-prediction averaging")
source = source.replace("approximately 638 of", "638 of")
source = source.replace("their best result reached 94% accuracy", "their ResNet50 deployment exceeded 90% accuracy")
section("This work develops a complete", "## II. Related Work", """
The contribution is an application-level integration and evaluation, not a new neural architecture, transfer-learning algorithm, or explanation method. EfficientNetB0 and Grad-CAM are established components [9], [10]. The specific contributions are:

1. A documented eight-condition classification task assembled from two public FDM sources, with harmonised labels and exclusion of ambiguous generic defect labels.
2. An auditable evaluation linking class-wise errors, sample support, confidence intervals, model size, and CPU inference cost to the retained deployed model.
3. A shared upload/live inference workflow combining temporal probability averaging, heuristic uncertainty warnings, curated defect guidance, and separately acquired printer telemetry.

The merged collection is a task-specific reuse of public data rather than a newly captured dataset. The study does not establish algorithmic superiority or claim that temporal averaging improves classification accuracy without a sequence-level experiment.
""")
source = source.replace("their best result reached 94% accuracy and was examined using Grad-CAM", "their ResNet50 deployment exceeded 90% accuracy on a small computing device and was examined using Grad-CAM")
section("For live use, OpenCV", "The dashboard maps", """
For live use, OpenCV captures camera images while inference is scheduled at a default interval of 0.8 s. The rolling window contains the five most recent successful prediction vectors, not five consecutive camera frames. For m available predictions (m <= 5), the displayed distribution is

> p̄(t) = [p(t) + ... + p(t − m + 1)] / m, m ≤ 5.

Five predictions were selected as a practical presentation heuristic; no frame-rate study, print-speed optimisation, or window-length ablation was performed. At the nominal 0.8 s prediction interval, the oldest and newest entries span about 3.2 s. If processing takes longer, the interval and span increase. This wall-clock span, rather than camera FPS, determines the smoothing delay. A full window replaces all pre-change entries after five new successful predictions. In an ideal abrupt change between two one-hot classes, the new class becomes dominant after three updates; a one-update transient contributes only 20% and may disappear from the displayed label. These are analytical properties, not measured physical failure-detection times. Therefore, short events can be suppressed and persistent failures can be reported late. There is no separately validated immediate-event bypass in the current software.

The dashboard flags a maximum displayed score below 0.60 as uncertain. This is an empirical warning threshold, not a calibrated 60% probability of correctness. Uploaded images use a single prediction without temporal averaging. Confidence calibration and threshold selection are discussed in Section V-B.

Grad-CAM remains an optional offline diagnostic from the project script. The live predictor returns class scores and guidance; it does not compute an explanation overlay for every video frame.
""")
section("### D. Explainability and Interface Behaviour", "## V. Discussion", """
### D. Interface Behaviour and Scope of Evidence

The upload and camera paths use the same retained model and class order. A still image returns an unsmoothed distribution; live predictions are averaged as specified in Section III-E. Temperatures and progress are controller measurements only when OctoPrint provides them. The CPU benchmark above uses synthetic input and excludes camera capture, decoding, network transport, and temporal decision delay. It establishes model-call feasibility, not end-to-end alarm latency or long-duration printer reliability. The optional Grad-CAM script supports offline inspection, but no quantitative explanation-quality study is available.

### E. Comparison with Published FDM Systems

Table V compares reported task scope and metrics from the cited studies. Values are transcribed from their publisher or author-institution records, rather than reproduced on the present split. NR denotes a value not established in those consulted records.

| Study / model | Task and data | Reported metric | Timing / scope |
|---|---|---|---|
| Kim et al. [4], ResNet50 | Spaghetti failure | Accuracy >90% | About 1 s; small device |
| Hu et al. [5], improved YOLOv8 | Five-defect detection | mAP50 97.5% | FPS +18.1% vs own baseline |
| Aktepe and Ergun [8], B0 | Five-class; 1,912 images | Accuracy 87.7% | Timing NR |
| Aktepe and Ergun [8], MobileNetV2 | Same source study | Accuracy 97.0% | Timing NR |
| Present EfficientNetB0 | Eight-class; 661 test images | Accuracy 96.52%; macro F1 96.95% | 108.3 ms; CPU model call |

The table provides context, not a controlled ranking: detection mAP50 cannot be equated with classification accuracy, and hardware, source data, class count, and splits differ. A same-split MobileNetV2/ResNet50 baseline and smoothing ablation have not been retained, so neither backbone superiority nor an accuracy gain from the interface can be concluded. The practical distinction is eight-condition coverage linked to class-wise error analysis and operator guidance.
""")
section("## V. Discussion", "## VI. Threats", """
## V. Discussion

### A. Error Patterns and Operator Decisions

The main weakness is the healthy/under-extrusion boundary: 12 of 135 no-defect images were assigned to under-extrusion. Better examples of borderline flow, illumination changes, and acceptable surface texture are a more targeted response than additional epochs alone. The 14-image off-platform test support also limits interpretation of perfect recall. Figure 2 makes these differences visible without hiding minority classes behind aggregate accuracy. Guidance is curated troubleshooting information; a class score does not establish a physical root cause or justify an automatic heater or axis adjustment.

### B. Calibration and the 60% Warning Rule

High classification accuracy does not establish calibrated confidence. Guo et al. [13] describe temperature scaling, which applies a positive scalar T to logits before softmax and fits T by minimising negative log-likelihood on a held-out calibration set. For this multiclass task, a single temperature is a simpler starting point than fitting separate Platt-style binary calibrators and then renormalising their outputs.

For M confidence bins, expected calibration error (ECE) compares the mean top-class confidence and empirical accuracy within each bin:

> ECE = Σ_b (n_b / n) |a_b − c_b|, b = 1, ..., M.

Here n_b is the number of samples in bin b, n is the total, a_b is bin accuracy and c_b is mean top-class confidence. A reproducible follow-up should report 15 equal-width bins, bin counts, reliability diagrams, negative log-likelihood, and the multiclass Brier score before and after scaling. T must be fitted using validation or a separate calibration subset, with test labels reserved for final reporting. A warning threshold should then be chosen on validation data according to acceptable false alarms and missed defects, reporting coverage and accuracy among accepted predictions. For video, calibration must also be assessed after probability averaging, since single-frame calibration does not establish calibration of the displayed stream.

The retained artefacts contain aggregate metrics, the confusion matrix, and model weights, but no paired validation/test probability arrays or original split manifest. ECE, a fitted temperature, and a threshold-performance curve therefore cannot be recovered from the available results alone. No such measurement is claimed here. The implemented 0.60 cutoff remains an unvalidated interface heuristic and must not be read as a statistical safety guarantee. Calibration, temporal-window ablation, and recording-grouped testing remain experimental requirements.
""")
section("## Reproducibility Notes", "## Acknowledgment", """
## Reproducibility Notes

The repository retains the model, class order, aggregate metrics, classification report, confusion matrix, training-curve image, and CPU benchmark. These support the reported image-level results. Exact calibration and grouped-split replication additionally require the selected image list, recording identifiers, and per-image probabilities, which are not retained in this package. The public source data must be obtained under their original conditions; a fresh download and random seed alone do not establish an identical historical split.
""")
source += '\n[13] C. Guo, G. Pleiss, Y. Sun, and K. Q. Weinberger, “On calibration of modern neural networks,” in *Proc. 34th International Conference on Machine Learning*, vol. 70, 2017, pp. 1321–1330. [Online]. Available: https://proceedings.mlr.press/v70/guo17a.html.\n'
# Use British English in prose while preserving exact published reference titles.
body, references = source.split("## References", 1)
for old, new in [("Optimization", "Optimisation"), ("optimization", "optimisation"),
                 ("localization", "localisation"), ("localizes", "localises"),
                 ("localize", "localise"), ("generalizability", "generalisability"),
                 ("normalized", "normalised"), ("visualization", "visualisation")]:
    body = body.replace(old, new)
source = body + "## References" + references
source = source.replace("The original labels *Under-Extrusion*", "The original labels *Under-Extrusion*")
manuscript = OUT / "IAME2026_Revised_Paper.md"
manuscript.write_text(source)

# Redraw both figures from retained integer counts and report metrics.
plt.rcParams.update({"font.family": "DejaVu Serif", "font.size": 10,
                     "pdf.fonttype": 42, "svg.fonttype": "none"})
names = ["Cracking", "Layer shifting", "No defect", "Off platform", "Spaghetti",
         "Stringing", "Under-extrusion", "Warping"]
cm = np.diag([68, 53, 121, 14, 36, 135, 131, 80])
for row, col, n in [(0,1,1),(0,7,2),(1,0,1),(2,6,12),(2,5,1),(2,4,1),
                    (6,2,2),(6,5,2),(7,3,1)]:
    cm[row, col] = n
assert cm.sum() == 661 and np.trace(cm) == 638
fig, ax = plt.subplots(figsize=(3.25, 3.25))
ax.imshow(cm, cmap="Blues", vmin=0, vmax=135)
short_names = ["Cracking", "Layer shift", "No defect", "Off platform", "Spaghetti", "Stringing", "Under-extr.", "Warping"]
ax.set_xticks(range(8), short_names, rotation=55, ha="right", fontsize=8)
ax.set_yticks(range(8), short_names, fontsize=8)
ax.set_xlabel("Predicted condition", fontsize=9)
ax.set_ylabel("Actual condition", fontsize=9)
for i in range(8):
    for j in range(8):
        ax.text(j, i, str(cm[i,j]), ha="center", va="center", fontsize=7,
                color="white" if cm[i,j] > 70 else "black")
fig.tight_layout()
for ext in ["png", "pdf", "svg"]:
    fig.savefig(OUT / f"figure1_confusion_matrix.{ext}", dpi=600)
plt.close(fig)
report = (ROOT / "outputs/reports/classification_report.txt").read_text()
rows = re.findall(r"^\s*(Cracking|Layer_shifting|No_defect|Off_platform|Spaghetti|Stringing|Under_extrusion|Warping)\s+([.\d]+)\s+([.\d]+)\s+([.\d]+)\s+(\d+)", report, re.M)
assert len(rows) == 8
fig, ax = plt.subplots(figsize=(3.25, 2.65))
y = np.arange(8)
for idx, (label, marker, color) in enumerate([("Precision", "o", "#146077"),
                                          ("Recall", "s", "#ad4b18"),
                                          ("F1-score", "^", "#39733b")]):
    ax.plot([float(r[idx+1])*100 for r in rows], y+(idx-1)*.2,
            linestyle="none", marker=marker, label=label, color=color, markersize=3.5)
ax.set_yticks(y, [f"{name} ({r[4]})" for name,r in zip(short_names,rows)], fontsize=8.5)
ax.invert_yaxis()
ax.set_xlim(85, 101)
ax.set_xticks([85, 90, 95, 100], ["85", "90", "95", "100"], fontsize=8.5)
ax.set_xlabel("Class-wise score (%)", fontsize=9)
ax.grid(axis="x", alpha=.3)
fig.legend(loc="lower center", ncol=3, frameon=False, fontsize=8,
           handletextpad=.3, columnspacing=.7)
fig.subplots_adjust(left=.40, right=.98, top=.98, bottom=.26)
for ext in ["png", "pdf", "svg"]:
    fig.savefig(OUT / f"figure2_class_performance.{ext}", dpi=600)
plt.close(fig)

spec = importlib.util.spec_from_file_location("paper_generator", ROOT / "tools/create_conference_paper.py")
generator = importlib.util.module_from_spec(spec)
spec.loader.exec_module(generator)
generator.SOURCE = manuscript
generator.IAME_OUTPUT = OUT / "IAME2026_Revised_Paper.docx"
generator.CONFUSION = OUT / "figure1_confusion_matrix.png"
generator.TRAINING_CURVES = OUT / "figure2_class_performance.png"
generator.FORCE_TABLE_COLUMNS = False
original_inline = generator.add_inline
def revised_inline(paragraph, text, size=9.5):
    return original_inline(paragraph, text, size=8.5 if size == 7.4 else size)
generator.add_inline = revised_inline
original_figure = generator.add_figure
def revised_figure(document, path, caption, width=3.2):
    if path == generator.TRAINING_CURVES:
        caption = "Fig. 2. Class-wise test performance. Parentheses show test support; Under-extr. denotes under-extrusion."
    return original_figure(document, path, caption, width=3.25)
generator.add_figure = revised_figure
original_table = generator.add_markdown_table
def revised_table(document, rows, number):
    original_table(document, rows, number)
    if rows[0].startswith("| Study / model"):
        # The comparison remains within a single template column.
        table = document.tables[-1]
        widths = [Inches(.95), Inches(.65), Inches(.80), Inches(.85)]
        from docx.oxml.ns import qn
        for col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
            col.set(qn("w:w"), str(int(width.inches * 1440)))
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = width
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(7.5)
generator.add_markdown_table = revised_table
generator.main()

response = """# Response to Reviewers - Paper 54

Title: A Real-Time Vision System for Multi-Class Defect Detection in Fused Deposition Modelling Using EfficientNetB0

The authors thank the reviewers for their constructive comments. The revised manuscript retains the measured eight-class results and distinguishes implemented features, analytical observations, and experiments that remain to be performed.

## Reviewer 1

### 1. Clarify novelty beyond EfficientNetB0, transfer learning and Grad-CAM

Section I now explicitly identifies the contribution as application-level integration and evaluation. No new architecture, transfer-learning algorithm, explanation method, or newly captured dataset is claimed. The contribution is the harmonised eight-condition task, auditable evaluation of the retained model, and a common upload/live workflow with guidance and controller-sourced telemetry. Section III-E clarifies that Grad-CAM is an optional offline diagnostic rather than a per-frame live feature.

### 2. Add a comparison table with existing methods and key metrics

Section IV-E and Table V compare Kim et al., Hu et al., and both models reported by Aktepe and Ergun with the present system. Task scope, reported metric, and available timing information are included. The values were checked against publisher/author-institution records. The table explicitly distinguishes mAP50 from classification accuracy, and different datasets and hardware prevent a controlled ranking. No unperformed baseline experiment is presented.

### 3. Shorten discussion and strengthen experimental comparison

The repeated general discussion and long reproducibility narrative have been shortened. The recovered space is used for Table V, explicit error interpretation, and the calibration discussion. The original 661-image test results remain unchanged. The revision identifies missing same-split backbone comparisons and temporal-window ablations instead of implying they were conducted.

## Reviewer 2

### 1. Improve Fig. 2 labels, legends and vector clarity

The previous wide raster training plot has been replaced by a class-performance plot generated from the retained classification report. It shows precision, recall and F1, with test support in each class label and a percentage axis. Both revised figures are supplied as vector PDF/SVG plus 600 dpi PNG fallbacks for Word. The final PDF embeds the vector plots. The original training image is retained in the project, but no numerical training history was invented to redraw it.

### 2. Explain the five-frame rolling window and transient failures

Section III-E corrects the description to five prediction vectors. The code schedules inference at a default 0.8 s interval, independent of camera FPS. A full window therefore spans approximately 3.2 s between oldest and newest entries. The choice is explicitly a presentation heuristic rather than an optimum derived from print speed. Analytical step-response behaviour and suppression of one-update transients are described; measured physical detection latency and an immediate-event bypass are not claimed.

### 3. Discuss ECE, Platt/temperature scaling and the 60% cutoff

Section V-B adds the ECE definition, a temperature-scaling procedure with validation-only fitting, and a proposed evaluation using reliability diagrams, negative log-likelihood, Brier score, coverage and accepted-prediction accuracy. Guo et al. is cited as [13]. Calibration of the averaged live output is distinguished from calibration of individual frames. The 60% cutoff is explicitly an unvalidated interface heuristic. The retained files lack paired validation/test prediction arrays and the original split manifest, so no ECE, fitted temperature, or statistically justified cutoff is fabricated. This comment is addressed through the requested discussion; empirical calibration remains outstanding.

### 4. Use consistent English

British English is used in the revised prose, including modelling, generalisation, normalisation, optimisation and localisation. Exact titles of published references retain their original spelling.

## Evidence and remaining experiments

The reported accuracy remains 638/661 = 96.52%; weighted F1 is 96.50% and macro F1 is 96.95%. Figure 1 is reconstructed from the retained confusion counts, which sum to 661 and contain 23 errors. Figure 2 uses the retained four-decimal class report. No retraining, new data collection, physical latency experiment, same-split backbone benchmark, calibration fit, or smoothing ablation was performed for this editorial revision.

Checked sources:
- Kim et al.: https://pure.konkuk.ac.kr/en/publications/systematic-deep-transfer-learning-method-based-on-a-small-image-d/
- Hu et al.: https://link.springer.com/article/10.1007/s00170-024-14452-4
- Aktepe and Ergun: https://dergipark.org.tr/en/pub/jmm/article/1910657
- Guo et al.: https://proceedings.mlr.press/v70/guo17a.html
"""
(OUT / "Response_to_Reviewers.md").write_text(response)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.2677), Inches(11.6929)
style = doc.styles["Normal"]
style.font.name, style.font.size = "Times New Roman", Pt(11)
for line in response.splitlines():
    if not line.strip():
        continue
    if line.startswith("### "):
        doc.add_heading(line[4:], 2)
    elif line.startswith("## "):
        doc.add_heading(line[3:], 1)
    elif line.startswith("# "):
        doc.add_heading(line[2:], 0)
    else:
        doc.add_paragraph(line)
doc.save(OUT / "Response_to_Reviewers.docx")
print(OUT)
