"""Build the final dated IAME manuscript and reviewer response from verified evidence."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import zipfile

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from sklearn.metrics import classification_report, log_loss
from docx import Document
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from docx.table import Table
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
OLD = ROOT / "outputs/paper/camera_ready_v1_2026-10-29"
REVIEW = ROOT / "outputs/paper/review_revision"
OUT = ROOT / "outputs/paper/IAME2026_Final"
OUT.mkdir(parents=True, exist_ok=True)
PAPER = "IAME2026_Paper54_Camera_Ready"
RESPONSE = "IAME2026_Paper54_Response_to_Reviewers"
RUN = ROOT / "outputs/reports/local_20261029_183228"


def evidence():
    saved = np.load(RUN / "validation_predictions.npz")
    y, probabilities = saved["labels"], saved["probabilities"].astype(float)
    probabilities /= probabilities.sum(axis=1, keepdims=True)
    names = json.loads((RUN / "metrics.json").read_text())["class_names"]
    report = classification_report(y, probabilities.argmax(1), labels=range(len(names)), target_names=names, output_dict=True, zero_division=0)
    confidence = probabilities.max(1)
    correct = probabilities.argmax(1) == y
    bins, ece = [], 0.0
    for index in range(15):
        mask = (confidence >= index/15) & ((confidence <= 1) if index == 14 else confidence < (index+1)/15)
        n = int(mask.sum())
        acc = float(correct[mask].mean()) if n else 0
        conf = float(confidence[mask].mean()) if n else 0
        ece += n / len(y) * abs(acc-conf)
        bins.append({"bin":index, "count":n,"accuracy":acc,"confidence":conf})
    summary = {"evaluation":"Reused grouped model-selection validation; not an independent test", "images":len(y),"accuracy":float(correct.mean()),"macro_f1":report["macro avg"]["f1-score"],"weighted_f1":report["weighted avg"]["f1-score"],"ece_15_equal_width":ece,"nll":float(log_loss(y,probabilities,labels=range(len(names)))),"multiclass_brier":float(np.mean(np.sum((probabilities-np.eye(len(names))[y])**2,axis=1))),"bins":bins,"per_class":{name:report[name] for name in names}}
    (OUT / "evaluation_evidence.json").write_text(json.dumps(summary,indent=2))
    return summary


def replace(paragraph, text):
    properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if properties is not None:
        run._r.insert(0, properties)


def build_manuscript(stats):
    if not (OLD / "IAME2026_Paper54_Camera_Ready_v1.docx").exists():
        if (OUT / f"{PAPER}.docx").exists():
            return Document(OUT / f"{PAPER}.docx")
        raise FileNotFoundError("Latest editable manuscript is missing")
    doc = Document(OLD / "IAME2026_Paper54_Camera_Ready_v1.docx")
    p = list(doc.paragraphs)
    updates = {
        3: "Abstract—Fused Deposition Modelling (FDM) can continue wasting material after a visible failure develops. This work integrates an eight-condition EfficientNetB0 classifier with a browser-based inspection workflow for remote operators. A historical 4,405-image experiment obtained 96.52% accuracy on a 661-image stratified test split; possible recording overlap limits that result. Subsequent offline continuation used 18,199 images, with 13,772 training and 4,427 recording/date-grouped validation images. The selected checkpoint reached 83.83% accuracy, 83.57% macro F1 and 83.75% weighted F1 on this reused model-selection validation set, not an independent test. A direct HTTPS interface keeps the camera preview local, analyses the newest frame without a server backlog, and presents stable decisions, uncertainty, curated causes and corrective checks. Photo upload, video playback analysis, snapshots, CSV export and on-demand Grad-CAM are integrated into a responsive light/dark dashboard. Face and image-quality checks plus operator view confirmation reduce accidental non-printer decisions but do not establish open-set recognition. Browser verification uses simulated media; external printer trials, calibrated alerts and machine control remain outside the demonstrated evidence.",
        8: "The system connects camera acquisition, classification and operator interpretation. A remote phone supplies live images to a host that runs the model. Class scores remain separate from physical process measurements: the application does not invent temperature, motion or nozzle values from an RGB image. The final interface reports recommendations only; printer control is not part of this submission.",
        11: "An auditable separation of historical image-level testing from later recording/date-grouped continuation, with retained split paths, probabilities and class-wise errors.",
        12: "A remote browser workflow with latest-frame processing, view rejection, consistent-reading alerts, Grad-CAM, troubleshooting guidance and downloadable inspection records.",
        21: "Two public sources define the eight-condition task [11], [12]. The first supplies cracking, layer shifting, off-platform, stringing and warping images; the second supplies good, under-extrusion, stringing and spaghetti frames with recording metadata. Labels are harmonised to one class order; Good becomes No defect, and generic Defected labels are excluded to avoid overlapping categories. The historical experiment used a capped 4,405-image collection (Table I). The restored continuation collection contains 18,199 images: 472 cracking, 364 layer shifting, 7,026 no defect, 91 off platform, 240 spaghetti, 3,956 stringing, 5,512 under-extrusion and 538 warping images. No corrupt image was found in the prepared collection.",
        22: "TABLE I\nHISTORICAL STRATIFIED IMAGE-LEVEL SPLIT",
        24: "The historical 70/15/15 stratified split used seed 42, producing 3,083/661/661 images. Its original split manifest is unavailable and recording overlap cannot be excluded. Continuation instead preserves Nimbus recording identifiers and groups original FDM filenames by capture date, keeping related copies together. The fixed split contains 13,772 training and 4,427 validation images. Grouping reduces within-recording leakage in this split but cannot remove unknown overlap with the source checkpoint's earlier training.",
        27: "Class imbalance is substantial, especially the 91 off-platform images. Historical training used balanced weights, computed as",
        29: "where N is training size, K the class count and n_c the class support. Local continuation used square-root balanced weights; the final experiment additionally multiplied cracking, no-defect, spaghetti and under-extrusion weights by 1.35. Weighting changes loss contributions without creating independent minority examples.",
        33: "Weighted sparse categorical cross-entropy is -w_y log(p_y). Historical training used Adam, batch size 32, initial head/backbone learning rates 0.001/0.00001, and the final 30 backbone layers for fine-tuning (Table II). Offline continuation on a GTX 1650 used batch size four and loss-based checkpoint selection. The final weak-class run completed two head epochs at 0.00001 and three backbone epochs at 0.000001 before early stopping. Batch-normalisation backbone layers remained frozen. Each continuation loaded prior weights with a new optimiser; this is not an exact interrupted-optimiser resume.",
        34: "TABLE II\nHISTORICAL NETWORK AND TRAINING CONFIGURATION",
        39: "Macro averaging gives each class equal importance, whereas weighted F1 reflects support. A Wilson 95% interval accompanies the historical image-level accuracy; it is not a cluster-aware confidence interval. Current grouped validation is reused for checkpoint selection and must not be presented as a new untouched test.",
        42: "The browser preview runs independently of inference. Each completed response schedules a fresh JPEG frame after a 35 ms scheduling floor; only one request per browser is in flight. Older camera frames are discarded rather than queued. Consequently, decision rate depends on host and network latency, not preview FPS. A 70% softmax-score threshold and three consecutive agreeing predictions produce the displayed confirmed condition. Other valid frames are labelled checking or uncertain; rejected views reset agreement history.",
        43: "Confirmation: c_t = c_(t-1) = c_(t-2), with min(s_t, s_(t-1), s_(t-2)) >= 0.70.",
        44: "The earlier local prototype averaged five prediction vectors at a nominal 0.8 s interval, spanning about 3.2 s from oldest to newest. That window was a presentation heuristic, not a result of printer-speed optimisation. The final browser version replaces averaging with agreement checks to avoid a fixed multi-second delay. Neither rule has a physical transient-failure or window-length ablation. Three agreeing updates need two inter-update intervals after the first qualifying result; a brief event may be missed. This is an analytical property, not a measured alarm guarantee.",
        45: "Before camera analysis, the operator confirms that the print is visible. Frontal/profile face checks and extreme lighting/near-uniform-view rejection run before classification. These are preliminary heuristics, not trained nozzle or bed detectors. A score threshold cannot guarantee rejection of arbitrary objects; all confirmed alerts still require visual review.",
        46: "Grad-CAM is available on demand for the current frame, photo or snapshot rather than every live update. This keeps explanation cost separate from continuous inference. The overlay visualises model attention and does not establish a defect boundary or causal explanation.",
        47: "The light/dark interface places confidence, possible causes and recommended checks beside the live view. Guidance is a fixed, class-conditioned engineering knowledge layer, not unrestricted generated advice. Photo and snapshot analysis use one unsmoothed prediction; video playback is inspected continuously at the inference rate, so not every encoded video frame is evaluated. CSV records include timestamp, source, video time, class score, displayed state, processing time and round-trip time. Camera tracks are released on Stop, mode changes, page hiding and page exit. No automatic heater, axis, cleaning or pause command is issued.",
        55: "The historical checkpoint correctly classified 638 of 661 held-out images: 96.52% accuracy, test loss 0.1058, macro F1 96.95% and weighted F1 96.50%. Its Wilson 95% image-level interval was 94.83-97.67%. Table IV and Figures 1-2 describe this historical experiment only, not the newly deployed continuation checkpoint.",
        60: "Historical stringing and spaghetti F1 scores were 98.90% and 98.63%. However, no-defect recall was 89.63%, with 12 acceptable images assigned to under-extrusion. This healthy/low-flow boundary is more consequential operationally than the aggregate score suggests.",
        61: "The 23 historical errors include three cracking, one layer-shifting, fourteen no-defect, four under-extrusion and one warping example. Correct recovery of every off-platform image involves only fourteen examples, giving a wide 78.47-100% Wilson recall interval. Such minority support does not establish dependable field performance.",
        65: "The historical model has 4,059,819 parameters and occupies 41.23 MB including its saved training state. Its thirty single-image CPU model calls, after five warm-ups, averaged 108.3 ms (median 107.6 ms) on the Ryzen 5 5600H laptop. These timings exclude capture, transport and alert confirmation. The selected continuation keeps the architecture but uses a different saved checkpoint; the historical timing is not assigned to the new network runtime.",
        67: "Automated browser checks exercise continuous acquisition, view confirmation, light/dark rendering, photo/video input, snapshots, Grad-CAM, CSV output and camera release. They use a synthetic camera and responsive browser viewports, not Android/iOS hardware or a physical printer. Network tests traverse a public HTTPS tunnel. These checks establish software-path operation, not physical defect sensitivity or cross-printer reliability. The reference machine-integration code is not enabled in the final public workflow.",
        74: "On the fixed 4,427-image grouped validation set, the historical checkpoint obtained 51.84% accuracy and 57.62% macro F1. Full-data continuation reached 83.10% accuracy and 84.98% macro F1; a short follow-up reached 83.49% and 83.42%. The final weak-class checkpoint achieved 83.83% accuracy, 83.57% macro F1 and 83.75% weighted F1, with loss 0.5254. It improved accuracy by fifteen images over the short run but did not improve every class. Spaghetti recall remains 58.33% (28/48); no-defect recall is 73.94%. This final checkpoint is selected for the demonstration, while all earlier weights are preserved. Validation was reused for model selection and source-checkpoint overlap is unknown: none of these figures is an independent test or directly comparable with the historical 96.52% result.",
        79: "Historical and current evaluations expose the same practical boundary: healthy surface variation can resemble weak extrusion. In grouped validation, spaghetti recall of 58.33% is a stronger warning than its 67.47% F1 alone. Three-update agreement may reduce visible flicker, but no alarm-accuracy improvement is claimed without labelled sequences. Curated advice suggests checks; it does not identify a verified physical root cause.",
        80: "B. Calibration and Uncertainty Thresholds",
        81: "Guo et al. [13] motivate post-hoc temperature scaling: a positive scalar fitted on calibration logits adjusts confidence without changing the top class. Expected calibration error (ECE) instead measures the absolute gap between confidence and empirical correctness across bins:",
        84: f"For the selected checkpoint, fifteen equal-width bins on the reused grouped validation probabilities give ECE {stats['ece_15_equal_width']:.4f}, negative log-likelihood {stats['nll']:.4f} and multiclass Brier score {stats['multiclass_brier']:.4f}. Bin counts and probabilities are retained. These are descriptive, uncalibrated model-selection diagnostics, not independent calibration evidence. Historical paired probabilities are unavailable.",
        85: "The earlier 60% warning and current 70% confirmation thresholds are interface heuristics, not statistically justified safety cutoffs. No fitted temperature, threshold optimum or before/after calibration gain is claimed. Calibration needs a separate subset and final recording-grouped or external testing; temporal decision rules also require sequence-level assessment.",
        87: "The historical image split may share recording-specific cues across subsets. Later grouping by Nimbus recording and FDM capture date is stronger, but its validation set was repeatedly inspected and its relation to earlier checkpoint training is unknown. Therefore the present record does not establish independent cross-printer generalisation. Small off-platform and spaghetti supports and mixed source definitions further limit precision.",
        88: "Whole-image classification cannot localise simultaneous defects. Face and lighting gates plus explicit operator confirmation reduce accidental background analysis, but arbitrary out-of-domain scenes can still pass. The confidence scores are uncalibrated. Synthetic browser checks are not long-duration physical trials, do not verify actual Android/iOS camera hardware and do not establish false-alarm or time-to-detection performance. Neither temperatures nor motion can be recovered reliably from a camera score.",
        89: "A follow-up should reserve an external printer-camera setup, record healthy and faulty sequences with timestamps, and measure sensitivity, false alarms and decision delay. Printer actuation must remain outside automatic control until confidence calibration and hardware safety checks are completed. The software can support this experiment, but missing physical evidence cannot be replaced by additional training epochs.",
        91: "An eight-condition FDM inspection system links EfficientNetB0 with continuous browser-camera analysis, images, videos, snapshots, on-demand Grad-CAM and operator feedback. Historical image-level testing yielded 96.52% accuracy; later grouped model-selection validation yielded 83.83% for the selected demonstration checkpoint. Keeping these results separate prevents an unsupported field-accuracy claim. The final application preserves a responsive preview while processing the latest frame without a backlog and exports auditable CSV records. Its contribution is an integrated decision-support workflow, not a new architecture or validated autonomous printer controller. External physical-printer and calibration evaluation remain necessary.",
        95: "The historical model, aggregate reports and project-derived figures are preserved separately from the selected continuation checkpoint, which retains split paths and per-image probabilities. Browser checks retain source, logs and screenshots. Source data remain subject to their original licences. No physical printer result, platform-wide similarity percentage or calibrated safety claim is asserted.",
    }
    for index,text in updates.items():
        replace(p[index], text)
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("[6]"):
            replace(paragraph,paragraph.text.replace("P. S. S. Sundari","P. S. Sundari"))
    for index in [62,63]:
        p[index]._element.getparent().remove(p[index]._element)
    # Add current per-class measurements without confusing them with historical Table IV.
    anchor = p[75]
    caption = deepcopy(p[56]._p)
    anchor._p.addprevious(caption)
    replace(Paragraph(caption, anchor._parent), "TABLE VI\nSELECTED CHECKPOINT: GROUPED VALIDATION")
    template = deepcopy(doc.tables[4]._tbl)
    anchor._p.addprevious(template)
    table = Table(template, anchor._parent)
    table.cell(0,0).text = "Class"
    table.cell(0,1).text = "Precision / Recall / F1"
    table.cell(0,2).text = "n"
    for row,(name,result) in zip(table.rows[1:],stats["per_class"].items()):
        row.cells[0].text = name.replace("_"," ")
        row.cells[1].text = f"{result['precision']:.4f} / {result['recall']:.4f} / {result['f1-score']:.4f}"
        row.cells[2].text = str(int(result["support"]))
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs: run.font.name="Times New Roman"; run.font.size=Pt(8)
    qa_path = ROOT / "logs/final_qa/validation.json"
    if qa_path.exists():
        checks = json.loads(qa_path.read_text())["checks"]
        sentence = " Local simulated-camera checks completed " + " and ".join(f"{c['updates_per_second']:.2f} updates/s in a {c['viewport']} viewport" for c in checks) + ". These short loopback checks are not physical camera, network-wide or failure-detection benchmarks."
        replace(p[67], updates[67] + sentence)
    public_path = ROOT / "logs/public_qa/validation.json"
    if public_path.exists():
        checks = json.loads(public_path.read_text())["checks"]
        rates = [c["updates_per_second"] for c in checks]
        replace(p[67],p[67].text + f" In corresponding short public-HTTPS simulated-camera runs, measured rates were {min(rates):.2f}-{max(rates):.2f} updates/s. These rates depend on the observed network and do not guarantee another device's throughput.")
    anchor=p[85]
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after=Pt(3)
    paragraph.add_run().add_picture(str(OUT/"figures/figure3_grouped_reliability.png"),width=Inches(3.22))
    anchor._p.addprevious(paragraph._p)
    caption=deepcopy(p[77]._p);anchor._p.addprevious(caption)
    replace(Paragraph(caption,anchor._parent),"Fig. 3. Raw reliability on reused grouped validation. The dashed diagonal denotes perfect calibration; no scaling was fitted.")
    return doc


def markdown(doc):
    lines=[]
    for element in doc.element.body.iterchildren():
        if element.tag.endswith("}p"):
            text=Paragraph(element,doc._body).text
            if text: lines.append(text)
        elif element.tag.endswith("}tbl"):
            rows=Table(element,doc._body).rows
            for index,row in enumerate(rows):
                lines.append("| "+" | ".join(c.text.replace("\n","; ") for c in row.cells)+" |")
                if index==0: lines.append("|"+"---|"*len(row.cells))
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def response_letter(stats):
    doc=Document()
    for section in doc.sections:
        section.page_width=Inches(8.2677); section.page_height=Inches(11.6929)
        section.top_margin=section.bottom_margin=Inches(.7)
        section.left_margin=section.right_margin=Inches(.8)
    doc.styles["Normal"].font.name="Calibri";doc.styles["Normal"].font.size=Pt(10.5)
    doc.add_heading("Response to Reviewers",0)
    doc.add_paragraph("IAME 2026 | Paper 54 | Final revision | 29 October 2026")
    doc.add_paragraph("A Real-Time Vision System for Multi-Class Defect Detection in Fused Deposition Modelling Using EfficientNetB0")
    items=[
        ("Reviewer 1.1 — Novel contribution", "The contribution is stated as application-level integration and evaluation, not a new EfficientNet, transfer-learning or Grad-CAM method. The manuscript now separates the historical experiment, grouped continuation and browser workflow, including limits of their evidence."),
        ("Reviewer 1.2 — Comparison with existing methods", "Table V compares cited FDM systems by task, reported metric and timing scope. Different datasets, class counts and mAP versus accuracy are explicitly not ranked. Section IV-F and Table VI add the completed 4,427-image grouped continuation results: the historical model's 51.84% accuracy on this split versus 83.83% for the selected checkpoint. Reused model-selection validation is not described as an independent test."),
        ("Reviewer 1.3 — Shorter discussion, stronger experiments", "The repetitive class discussion is shortened. The saved grouped split, per-class results, continuation history and calibration diagnostics now occupy the recovered space. The original 96.52% historical score is retained but is not used as the new deployment score."),
        ("Reviewer 2.1 — Figure clarity", "All three project-derived manuscript figures are embedded as vector PDF content. Class labels and legends remain searchable. Figures 1-2 are labelled as historical results; Table VI and the new reliability diagram in Figure 3 report the grouped validation. No online photographs or third-party plots are included."),
        ("Reviewer 2.2 — Five-frame window and transients", "Section III-E explains that the earlier five-prediction window was heuristic, not selected from printer speed or camera FPS. Its nominal span was 3.2 s at 0.8 s per update. The final browser workflow instead processes the newest frame after each response and requires three agreeing predictions at 70% confidence. Short events can still be missed, and no physical failure latency or temporal ablation is claimed."),
        ("Reviewer 2.3 — Calibration and the uncertainty threshold", f"Section V-B discusses temperature scaling and reports descriptive raw calibration diagnostics on the saved grouped validation probabilities: 15-bin ECE {stats['ece_15_equal_width']:.4f}, NLL {stats['nll']:.4f}, multiclass Brier {stats['multiclass_brier']:.4f}. Validation reuse and unknown source-training overlap are stated. No independent calibration assessment, fitted temperature or optimised threshold is claimed. Historical paired probabilities remain unavailable."),
        ("Reviewer 2.4 — Consistent English", "British spelling is used in the narrative: modelling, normalisation, optimisation, generalisation and labelled. Published reference titles retain their original spelling. Author names use trailing initials; Assistant Professor appears on a separate, non-bold line for the two faculty authors."),
        ("Evidence and submission boundary", "Browser verification uses simulated media and responsive Chromium viewports, not a physical printer or actual Android/iOS hardware. Printer actuation is not enabled. The acceptance email specifies camera-ready completion by 31 October 2026 and similarity below 15%; the author's stricter under-10% target still requires an institution-approved Turnitin/iThenticate check. No such report is fabricated or included. Supervisor approval, registration and CMT upload remain author actions."),
    ]
    for heading,text in items:
        doc.add_heading(heading,2);doc.add_paragraph(text)
    doc.save(OUT/f"{RESPONSE}.docx")
    (OUT/f"{RESPONSE}.md").write_text(markdown(doc))


def main():
    stats=evidence()
    figures=OUT/"figures"; figures.mkdir(exist_ok=True)
    populated=[b for b in stats["bins"] if b["count"]]
    plt.rcParams.update({"font.size":8,"axes.labelsize":8.5,"xtick.labelsize":8,"ytick.labelsize":8})
    fig,ax=plt.subplots(figsize=(3.22,2.45))
    ax.plot([0,1],[0,1],"--",color="#7c858c",linewidth=1,label="Perfect calibration")
    ax.plot([b["confidence"] for b in populated],[b["accuracy"] for b in populated],"o-",color="#087f78",label="Selected checkpoint",linewidth=1.5,markersize=4)
    ax.set(xlim=(0,1),ylim=(0,1),xlabel="Mean predicted confidence",ylabel="Observed accuracy")
    ax.grid(alpha=.18);ax.legend(fontsize=7.5,loc="upper left")
    fig.tight_layout()
    fig.savefig(figures/"figure3_grouped_reliability.png",dpi=400)
    fig.savefig(figures/"figure3_grouped_reliability.pdf")
    fig.savefig(figures/"figure3_grouped_reliability.svg")
    svg=figures/"figure3_grouped_reliability.svg"
    svg.write_text("\n".join(line.rstrip() for line in svg.read_text().splitlines())+"\n")
    plt.close(fig)
    doc=build_manuscript(stats)
    reliability=doc.inline_shapes[2]
    relationship=reliability._inline.graphic.graphicData.pic.blipFill.blip.embed
    doc.part.related_parts[relationship]._blob=(figures/"figure3_grouped_reliability.png").read_bytes()
    reliability.height=Inches(2.45)
    doc.save(OUT/f"{PAPER}.docx")
    (OUT/f"{PAPER}.md").write_text(markdown(doc))
    response_letter(stats)
    for stem in ["figure1_confusion_matrix","figure2_class_performance"]:
        for ext in ["pdf","svg","png"]:
            source=REVIEW/f"{stem}.{ext}"
            if source.exists(): shutil.copy2(source,figures/source.name)
    profile=ROOT/".tools/libreoffice-final"
    subprocess.run(["libreoffice",f"-env:UserInstallation={profile.as_uri()}","--headless","--convert-to","pdf","--outdir",str(OUT),str(OUT/f"{PAPER}.docx"),str(OUT/f"{RESPONSE}.docx")],check=True,timeout=120)
    pdfpath=OUT/f"{PAPER}.pdf"
    with pymupdf.open(pdfpath) as pdf:
        assert len(pdf)<=6, f"Paper exceeds six pages: {len(pdf)}"
        count=0
        for page in pdf:
            for info in page.get_image_info(xrefs=True):
                if info["width"]<500:continue
                name=["figure1_confusion_matrix.pdf","figure2_class_performance.pdf","figure3_grouped_reliability.pdf"][count]
                rect=pymupdf.Rect(info["bbox"])
                with pymupdf.open(figures/name) as vector:
                    page.add_redact_annot(rect,fill=(1,1,1));page.apply_redactions(images=2,graphics=0,text=1)
                    page.show_pdf_page(rect,vector,0)
                count+=1
        assert count==3, count
        temp=OUT/"vector_export.pdf";pdf.save(temp,garbage=4,deflate=True)
    temp.replace(pdfpath)
    shutil.copy2(pdfpath,OUT/"IAME2026_Paper54_Full_Paper.pdf")
    checks={}
    for stem in [PAPER,RESPONSE]:
        with pymupdf.open(OUT/f"{stem}.pdf") as pdf:
            text="\n".join(page.get_text() for page in pdf)
            assert not re.search(r"\bwe\b",text,re.I)
            for value in ["83.83%","96.52%","4,427"]:assert value in text,(stem,value)
            assert not pdf.is_encrypted
            assert all(abs(page.rect.width-595.28)<2 and abs(page.rect.height-841.89)<2 for page in pdf)
            checks[stem]={"pages":len(pdf),"size":"A4","sha256":hashlib.sha256((OUT/f"{stem}.pdf").read_bytes()).hexdigest()}
            for index,page in enumerate(pdf):page.get_pixmap(matrix=pymupdf.Matrix(1,1)).save(OUT/f"{stem}_preview_{index+1}.png")
            (OUT/f"{stem}_extracted.txt").write_text("\n".join(line.rstrip() for line in text.splitlines()).strip() + "\n")
    (OUT/"validation.json").write_text(json.dumps(checks,indent=2))
    (OUT/"READ_ME.txt").write_text("IAME 2026 | Paper 54 | Final package | 29 October 2026\n\nUpload IAME2026_Paper54_Camera_Ready.pdf as the manuscript.\nIAME2026_Paper54_Full_Paper.pdf is the same manuscript under an alternate name, not an additional paper.\nUpload the reviewer response only if CMT requests it. DOCX and Markdown are editable.\n\nCurrent checkpoint: 83.83% on reused grouped validation, not an independent test.\nHistorical checkpoint: 96.52% on the earlier image-level test. Do not conflate the two.\n\nThe acceptance email specifies 31 October 2026 and similarity below 15%.\nUnder-10% similarity cannot be verified without an external similarity report.\nNo physical printer, automatic machine control, actual iOS/Android hardware proof, or formal explanation-quality study is claimed.\nSupervisors must approve the final manuscript and institution-approved similarity report before CMT submission.\nNothing has been submitted, registered or paid on the author's behalf.\n")
    with zipfile.ZipFile(OUT/"IAME2026_Paper54_Submission_Package.zip","w",zipfile.ZIP_DEFLATED) as package:
        for path in sorted(OUT.iterdir()):
            if path.is_file() and (path.suffix in {".pdf",".docx",".md",".json"} or path.name=="READ_ME.txt"):package.write(path,path.name)
        for path in sorted(figures.iterdir()):package.write(path,"figures/"+path.name)
    print(json.dumps(checks,indent=2))
    print(OUT)


if __name__=="__main__":main()
